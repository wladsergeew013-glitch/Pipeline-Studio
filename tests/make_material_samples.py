from pathlib import Path
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas
from PIL import Image

dest = Path(__file__).resolve().parents[1] / 'qa' / 'material-samples'
dest.mkdir(parents=True,exist_ok=True)
doc = Document()
doc.add_heading('Pipeline Studio material test', 0)
doc.add_paragraph('Synthetic Word document for opening and relative path checks.')
doc.save(dest / 'Sample.docx')
deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[1])
slide.shapes.title.text = 'Pipeline Studio material test'
slide.placeholders[1].text = 'Synthetic presentation with one slide.'
deck.save(dest / 'Sample.pptx')
pdf = canvas.Canvas(str(dest / 'Preview.pdf'))
pdf.drawString(60, 780, 'Pipeline Studio synthetic PDF preview')
pdf.save()
Image.new('RGB', (160, 100), '#d5e8d4').save(dest / 'Sample.png')
